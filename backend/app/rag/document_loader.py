import os
import io
import re
import pandas as pd
from pathlib import Path
from typing import List, Dict, Any, Tuple
from langchain_core.documents import Document

class DocumentLoaderError(Exception):
    pass

class DocumentLoader:
    """Production-grade document loader supporting PDF, DOCX, TXT, CSV, and Markdown with clean text extraction."""

    @staticmethod
    def clean_text(text: str) -> str:
        """Cleans extracted text by normalizing unicode characters, spaces, and weird line breaks."""
        if not text:
            return ""
        # Remove null bytes
        text = text.replace("\x00", " ")
        # Normalize bullet points and quotes to standard characters
        text = text.replace("\u25cf", "• ").replace("\u2022", "• ").replace("\u2013", "-").replace("\u2014", "-")
        text = text.replace("\u2018", "'").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
        # Normalize multiple spaces and tabs
        text = re.sub(r"[ \t]+", " ", text)
        # Normalize excessive newlines
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @classmethod
    def load_file(cls, file_path: str, original_filename: str = None) -> Tuple[List[Document], Dict[str, Any]]:
        """
        Loads and parses a document into a list of LangChain Document objects with page tracking.
        Returns (list_of_documents, document_metadata)
        """
        path = Path(file_path)
        if not path.exists():
            raise DocumentLoaderError(f"File not found: {file_path}")

        filename = original_filename or path.name
        ext = path.suffix.lower()
        file_size = path.stat().st_size

        if file_size == 0:
            raise DocumentLoaderError(f"File '{filename}' is empty (0 bytes).")

        docs: List[Document] = []
        extra_meta: Dict[str, Any] = {
            "original_name": filename,
            "file_type": ext,
            "file_size": file_size,
            "total_pages": 1,
            "columns": None
        }

        try:
            if ext == ".pdf":
                docs, total_pages = cls._load_pdf(path, filename)
                extra_meta["total_pages"] = total_pages

            elif ext in [".docx", ".doc"]:
                docs = cls._load_docx(path, filename)
                extra_meta["total_pages"] = 1

            elif ext in [".txt", ".md"]:
                docs = cls._load_txt_or_md(path, filename, ext)
                extra_meta["total_pages"] = 1

            elif ext == ".csv":
                docs, columns = cls._load_csv(path, filename)
                extra_meta["total_pages"] = 1
                extra_meta["columns"] = columns

            else:
                raise DocumentLoaderError(f"Unsupported file format '{ext}'. Supported: .pdf, .docx, .txt, .csv, .md")

        except Exception as e:
            if isinstance(e, DocumentLoaderError):
                raise e
            raise DocumentLoaderError(f"Failed to parse document '{filename}': {str(e)}")

        if not docs or all(len(d.page_content.strip()) == 0 for d in docs):
            raise DocumentLoaderError(f"Document '{filename}' produced no readable text content.")

        return docs, extra_meta

    @classmethod
    def _load_pdf(cls, path: Path, filename: str) -> Tuple[List[Document], int]:
        from pypdf import PdfReader
        
        try:
            reader = PdfReader(str(path))
            total_pages = len(reader.pages)
            if total_pages == 0:
                raise DocumentLoaderError("PDF contains 0 pages.")

            docs = []
            for page_idx, page in enumerate(reader.pages, start=1):
                raw_text = page.extract_text() or ""
                cleaned = cls.clean_text(raw_text)
                if cleaned:
                    docs.append(Document(
                        page_content=cleaned,
                        metadata={
                            "document_name": filename,
                            "page_number": page_idx,
                            "file_type": ".pdf",
                            "is_tabular": False
                        }
                    ))
            
            if not docs:
                docs.append(Document(
                    page_content=f"[Scanned or empty PDF document: {filename}]",
                    metadata={
                        "document_name": filename,
                        "page_number": 1,
                        "file_type": ".pdf",
                        "is_tabular": False
                    }
                ))
            return docs, total_pages
        except Exception as e:
            raise DocumentLoaderError(f"Invalid or corrupted PDF file: {str(e)}")

    @classmethod
    def _load_docx(cls, path: Path, filename: str) -> List[Document]:
        import docx
        try:
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            cleaned = cls.clean_text(full_text)
            
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        table_texts.append(row_text)
            
            if table_texts:
                cleaned += "\n\n### Tables:\n" + "\n".join(table_texts)

            return [Document(
                page_content=cleaned,
                metadata={
                    "document_name": filename,
                    "page_number": 1,
                    "file_type": ".docx",
                    "is_tabular": False
                }
            )]
        except Exception as e:
            raise DocumentLoaderError(f"Invalid or corrupted DOCX file: {str(e)}")

    @classmethod
    def _load_txt_or_md(cls, path: Path, filename: str, ext: str) -> List[Document]:
        try:
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = path.read_text(encoding="latin-1")

            cleaned = cls.clean_text(content)
            return [Document(
                page_content=cleaned,
                metadata={
                    "document_name": filename,
                    "page_number": 1,
                    "file_type": ext,
                    "is_tabular": False
                }
            )]
        except Exception as e:
            raise DocumentLoaderError(f"Error reading text file: {str(e)}")

    @classmethod
    def _load_csv(cls, path: Path, filename: str) -> Tuple[List[Document], List[str]]:
        """
        Loads CSV records cleanly without dumping raw df.describe() statistical matrices into document chunks.
        """
        try:
            try:
                df = pd.read_csv(str(path))
            except UnicodeDecodeError:
                df = pd.read_csv(str(path), encoding="latin-1")

            columns = [str(c) for c in df.columns.tolist()]
            row_count = len(df)
            
            # Format row records cleanly as structured text
            row_summaries = []
            for idx, row in df.iterrows():
                row_items = [f"{col}: {val}" for col, val in row.items() if pd.notna(val)]
                row_str = f"Data Record #{idx+1} ({filename}): " + ", ".join(row_items)
                row_summaries.append(row_str)
                if len(row_summaries) >= 100:
                    break

            content_text = f"Dataset: {filename} ({row_count} total records)\nColumns: {', '.join(columns)}\n\n" + "\n".join(row_summaries)
            
            doc = Document(
                page_content=content_text,
                metadata={
                    "document_name": filename,
                    "page_number": 1,
                    "file_type": ".csv",
                    "is_tabular": True,
                    "total_rows": row_count,
                    "columns": columns
                }
            )
            return [doc], columns
        except Exception as e:
            raise DocumentLoaderError(f"Error parsing CSV file: {str(e)}")
